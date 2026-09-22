import io
import pytest
from datetime import datetime
from openpyxl import load_workbook
from app.models.registration import Registration
from app.models.payment import Payment


def test_admin_can_view_registrations(client, admin_user, student_user, sample_visit, db):
    """Test admin view of student registrations."""
    client.post('/auth/login', data={'email': admin_user.email, 'password': 'Admin@123'})

    reg = Registration(
        registration_id="IV-2026-123456",
        user_id=student_user.id,
        visit_id=sample_visit.id,
        status="confirmed",
        emergency_contact="Parent",
        emergency_phone="9811122233",
        terms_accepted=True
    )
    db.session.add(reg)
    db.session.commit()

    response = client.get('/admin/registrations')
    assert response.status_code == 200
    assert b'IV-2026-123456' in response.data
    assert student_user.full_name.encode() in response.data


def test_admin_can_view_payments(client, admin_user, student_user, sample_visit, db):
    """Test admin financial payments ledger."""
    client.post('/auth/login', data={'email': admin_user.email, 'password': 'Admin@123'})

    reg = Registration(
        registration_id="IV-2026-654321",
        user_id=student_user.id,
        visit_id=sample_visit.id,
        status="confirmed",
        emergency_contact="Parent",
        emergency_phone="9811122233",
        terms_accepted=True
    )
    db.session.add(reg)
    db.session.flush()

    payment = Payment(
        registration_id=reg.id,
        razorpay_payment_id="pay_test_7788",
        amount=sample_visit.fee,
        currency="INR",
        payment_method="DevMock UPI",
        status="successful",
        paid_at=datetime.utcnow()
    )
    db.session.add(payment)
    db.session.commit()

    response = client.get('/admin/payments')
    assert response.status_code == 200
    assert b'pay_test_7788' in response.data
    assert b'500.00' in response.data


def test_registration_export_csv(client, admin_user, student_user, sample_visit, db):
    """Test CSV export generation."""
    client.post('/auth/login', data={'email': admin_user.email, 'password': 'Admin@123'})

    reg = Registration(
        registration_id="IV-2026-CSV01",
        user_id=student_user.id,
        visit_id=sample_visit.id,
        status="confirmed",
        emergency_contact="Parent",
        emergency_phone="9811122233",
        terms_accepted=True
    )
    db.session.add(reg)
    db.session.commit()

    response = client.get('/admin/registrations/export/csv')
    assert response.status_code == 200
    assert response.mimetype == 'text/csv'
    assert b'Registration ID' in response.data
    assert b'IV-2026-CSV01' in response.data
    assert student_user.full_name.encode() in response.data


def test_registration_export_excel(client, admin_user, student_user, sample_visit, db):
    """Test Excel (.xlsx) export file validity."""
    client.post('/auth/login', data={'email': admin_user.email, 'password': 'Admin@123'})

    reg = Registration(
        registration_id="IV-2026-XLSX01",
        user_id=student_user.id,
        visit_id=sample_visit.id,
        status="confirmed",
        emergency_contact="Parent",
        emergency_phone="9811122233",
        terms_accepted=True
    )
    db.session.add(reg)
    db.session.commit()

    response = client.get('/admin/registrations/export/excel')
    assert response.status_code == 200
    assert 'spreadsheetml' in response.mimetype

    # Verify that the returned stream is a valid openpyxl workbook
    wb = load_workbook(filename=io.BytesIO(response.data))
    ws = wb.active
    assert ws.title == "Registrations"
    # Find cell with registration ID
    found = False
    for row in ws.iter_rows(values_only=True):
        if "IV-2026-XLSX01" in row:
            found = True
            break
    assert found is True


def test_dashboard_statistics(client, admin_user, student_user, sample_visit, db):
    """Verify admin dashboard calculates dynamic real data and returns HTTP 200."""
    client.post('/auth/login', data={'email': admin_user.email, 'password': 'Admin@123'})

    reg = Registration(
        registration_id="IV-2026-STAT01",
        user_id=student_user.id,
        visit_id=sample_visit.id,
        status="confirmed",
        emergency_contact="Parent",
        emergency_phone="9811122233",
        terms_accepted=True
    )
    db.session.add(reg)
    db.session.flush()

    payment = Payment(
        registration_id=reg.id,
        razorpay_payment_id="pay_stat_1",
        amount=500.0,
        currency="INR",
        status="successful",
        paid_at=datetime.utcnow()
    )
    db.session.add(payment)
    db.session.commit()

    response = client.get('/admin/dashboard')
    assert response.status_code == 200
    assert b'Total Students' in response.data
    assert b'Total Revenue Collected' in response.data
    assert b'Tata Motors' in response.data


def test_admin_reports_and_students_directory(client, admin_user, student_user, sample_visit, db):
    """Test admin reports hub and students management."""
    client.post('/auth/login', data={'email': admin_user.email, 'password': 'Admin@123'})

    # Reports
    rep_resp = client.get('/admin/reports')
    assert rep_resp.status_code == 200
    assert b'Utilization &amp; Revenue' in rep_resp.data or b'Utilization & Revenue' in rep_resp.data

    # Students Directory
    std_resp = client.get('/admin/students')
    assert std_resp.status_code == 200
    assert student_user.full_name.encode() in std_resp.data

    # Toggle Student Status
    toggle_resp = client.get(f'/admin/students/{student_user.id}/toggle-status', follow_redirects=True)
    assert toggle_resp.status_code == 200
    assert student_user.is_active is False
