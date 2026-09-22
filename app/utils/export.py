import io
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def generate_registrations_csv(registrations) -> str:
    """Generates CSV string from registration records."""
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Header row
    writer.writerow([
        'Registration ID',
        'Student Name',
        'College Email',
        'Roll Number',
        'Department',
        'Academic Year',
        'Division',
        'Phone Number',
        'Visit Title',
        'Company Name',
        'Visit Date',
        'Visit Location',
        'Registration Status',
        'Registered On',
        'Amount Paid (INR)',
        'Payment Status',
        'Payment ID',
        'Emergency Contact',
        'Emergency Phone'
    ])
    
    for reg in registrations:
        user = reg.student
        visit = reg.visit
        payment = reg.latest_payment
        
        writer.writerow([
            reg.registration_id,
            user.full_name if user else 'N/A',
            user.email if user else 'N/A',
            user.roll_number if user else 'N/A',
            user.department if user else 'N/A',
            user.year if user else 'N/A',
            user.division if user else 'N/A',
            user.phone if user else 'N/A',
            visit.title if visit else 'N/A',
            visit.company_name if visit else 'N/A',
            visit.visit_date.strftime('%Y-%m-%d') if visit and visit.visit_date else 'N/A',
            visit.location if visit else 'N/A',
            reg.status.upper(),
            reg.registered_at.strftime('%Y-%m-%d %H:%M:%S') if reg.registered_at else 'N/A',
            f"{payment.amount:.2f}" if payment else '0.00',
            payment.status.upper() if payment else 'UNPAID',
            payment.razorpay_payment_id if payment and payment.razorpay_payment_id else 'N/A',
            reg.emergency_contact,
            reg.emergency_phone
        ])
        
    return output.getvalue()


def generate_registrations_excel(registrations) -> io.BytesIO:
    """Generates styled Excel workbook (.xlsx) as a BytesIO stream."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Registrations"
    
    # Title Banner
    ws.merge_cells('A1:S1')
    title_cell = ws['A1']
    title_cell.value = "College Industrial Visit Portal — Registration Ledger"
    title_cell.font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 40
    
    # Subtitle / Export timestamp
    ws.merge_cells('A2:S2')
    sub_cell = ws['A2']
    sub_cell.value = f"Exported on: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')} | Total Records: {len(registrations)}"
    sub_cell.font = Font(name='Calibri', size=11, italic=True, color='374151')
    sub_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 20
    
    # Table Headers
    headers = [
        'Registration ID',
        'Student Name',
        'College Email',
        'Roll Number',
        'Department',
        'Academic Year',
        'Division',
        'Phone Number',
        'Visit Title',
        'Company Name',
        'Visit Date',
        'Location',
        'Registration Status',
        'Registered At',
        'Amount Paid (₹)',
        'Payment Status',
        'Payment ID',
        'Emergency Contact',
        'Emergency Phone'
    ]
    
    ws.append([])  # Blank row 3
    ws.append(headers)  # Row 4
    ws.row_dimensions[4].height = 28
    
    header_fill = PatternFill(start_color='3B82F6', end_color='3B82F6', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
        
    # Populate Data
    for reg in registrations:
        user = reg.student
        visit = reg.visit
        payment = reg.latest_payment
        
        row_data = [
            reg.registration_id,
            user.full_name if user else 'N/A',
            user.email if user else 'N/A',
            user.roll_number if user else 'N/A',
            user.department if user else 'N/A',
            user.year if user else 'N/A',
            user.division if user else 'N/A',
            user.phone if user else 'N/A',
            visit.title if visit else 'N/A',
            visit.company_name if visit else 'N/A',
            visit.visit_date.strftime('%Y-%m-%d') if visit and visit.visit_date else 'N/A',
            visit.location if visit else 'N/A',
            reg.status.upper(),
            reg.registered_at.strftime('%Y-%m-%d %H:%M') if reg.registered_at else 'N/A',
            payment.amount if payment else 0.0,
            payment.status.upper() if payment else 'UNPAID',
            payment.razorpay_payment_id if payment and payment.razorpay_payment_id else 'N/A',
            reg.emergency_contact,
            reg.emergency_phone
        ]
        ws.append(row_data)
        
    # Style rows
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = Font(name='Calibri', size=10)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')
            
    # Auto-adjust column widths
    from openpyxl.utils import get_column_letter
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
